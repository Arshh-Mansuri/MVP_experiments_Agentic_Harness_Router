"""Build the success table the table router puts in its prompt, then freeze it with a hash.

Table A: GPT-5.6-Luna Phase 1 results. Development tasks by default; --include-test adds the 18 sealed test
         tasks so the lookup covers all 45 tasks Luna has run, which is a deliberate, logged unsealing that
         turns those tasks into training data (see PROBLEM_DEFINITION.md).
Table B: Qwen3-Coder 480B baseline on all 89 Terminal-Bench tasks (Qwen_89_Task_Comparison_2026-09-20.xlsx),
         one run per task and harness; a weak hint from a different model.

The best harness per task is the stable oracle of study/analyze_phase1.py, reused rather than reimplemented so
the table and the analysis can never disagree: highest mean pass rate, ties to the cheaper harness, and only
harnesses that actually produced a scored run.

usage: python3 router/build_success_table.py [--include-test]
writes router/success_table.json, router/success_table.md and the table router's router/table_system_prompt.txt
"""
import argparse, datetime, hashlib, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(ROOT, "study"))
import analyze_phase1 as an
import profile_router as pr

HARNESSES = pr.HARNESSES
DEFAULT = "mini-swe-agent"
QWEN_XLSX = os.path.join(ROOT, "archive", "Qwen_89_Task_Comparison_2026-09-20.xlsx")
QWEN_COLS = {"mini-SWE": "mini-swe-agent", "Terminus-2": "terminus-2", "OpenCode": "opencode", "pi": "pi"}
OUT_JSON = os.path.join(HERE, "success_table.json")
OUT_MD = os.path.join(HERE, "success_table.md")


def first_sentence(task):
    try:
        text = pr.task_text(task)
    except FileNotFoundError:
        return ""
    text = re.sub(r"\s+", " ", re.sub(r"[#*`>]", "", text)).strip()
    m = re.match(r"(.{20,200}?[.!?])(\s|$)", text)
    return (m.group(1) if m else text[:200]).strip()


def table_a(tasks, split_of, cells):
    rows = []
    for task in sorted(tasks):
        stats = {}
        for h in HARNESSES:
            xs = cells[task][h]
            if xs:
                c = an.cell(cells, task, h)
                stats[h] = {"passes": int(sum(x["reward"] for x in xs)), "runs": len(xs),
                            "avg_cost": round(c["cost"], 4)}
        ran = an.ran(cells, task)
        if ran:
            best = an.stable_pick(cells, task)
            rate = lambda h: stats[h]["passes"] / stats[h]["runs"]
            notes = []
            if stats[best]["passes"] == 0:
                notes.append("no harness solved it; cheapest that ran chosen")
            elif len([h for h in ran if rate(h) == rate(best)]) > 1:
                notes.append("tie on pass rate; cheapest chosen")
            missing = [h for h in HARNESSES if h not in ran]
            if missing:
                notes.append(f"never produced a result: {', '.join(missing)}")
            note = "; ".join(notes)
        else:
            best, note = DEFAULT, "no harness produced a result; default"
        rows.append({"task": task, "split": split_of[task], "about": first_sentence(task),
                     "results": stats, "best": best, "note": note})
    return rows


def table_b():
    from openpyxl import load_workbook
    wb = load_workbook(QWEN_XLSX, read_only=True, data_only=True)
    rows, header = [], None
    for r in wb["Results"].iter_rows(values_only=True):
        if r and r[0] == "Task":
            header = list(r)
            continue
        if not header or not r or not r[0]:
            if header and rows:
                break
            continue
        res = {}
        for col, h in QWEN_COLS.items():
            cell = r[header.index(col)]
            res[h] = str(cell).split("\n")[0].strip() if cell else "n/a"
        rows.append({"task": r[0], "results": res})
    return rows


def render_md(t):
    scope = ("Luna Phase 1 development and test tasks (test unsealed deliberately: these tasks are now training "
             "data for the lookup, not a held-out set)" if t["includes_test"]
             else "Luna Phase 1 development tasks only")
    out = ["# Success table (frozen)", "",
           f"sha256 {t['table_sha256']}, built {t['built']}. Table A = {scope}.", "",
           "## Table A: GPT-5.6-Luna, up to 3 runs per harness (passes/runs, average cost per run)", "",
           "| task | split | " + " | ".join(HARNESSES) + " | best | note |", "|" + "---|" * (len(HARNESSES) + 4)]
    for r in t["table_a"]:
        cells = [f"{r['results'][h]['passes']}/{r['results'][h]['runs']} ${r['results'][h]['avg_cost']:.3f}"
                 if h in r["results"] else "-" for h in HARNESSES]
        out.append(f"| {r['task']} | {r['split']} | " + " | ".join(cells) + f" | **{r['best']}** | {r['note']} |")
    out += ["", "## Table B: Qwen3-Coder 480B, 1 run per harness (weak hint; OpenCode not available to us)", "",
            "| task | mini-swe-agent | terminus-2 | pi | opencode |", "|---|---|---|---|---|"]
    for r in t["table_b"]:
        res = r["results"]
        out.append(f"| {r['task']} | {res['mini-swe-agent']} | {res['terminus-2']} | {res['pi']} | {res['opencode']} |")
    return "\n".join(out) + "\n"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--include-test", action="store_true",
                    help="add the 18 sealed test tasks to Table A (logged to study/test_unseal_log.jsonl); "
                         "they become training data and can no longer measure generalisation")
    a = ap.parse_args()
    split = json.load(open(os.path.join(ROOT, "study", "phase1_tasks.json")))
    tasks = split["dev"] + split["test"] if a.include_test else split["dev"]
    split_of = {t: "dev" for t in split["dev"]}
    split_of.update({t: "test" for t in split["test"]})

    if a.include_test:
        rec = {"ts": datetime.datetime.now().isoformat(timespec="seconds"), "by": "build_success_table.py",
               "purpose": "hard-code the best harness for all 45 Luna tasks in Table A (deployment lookup)",
               "test_tasks": len(split["test"])}
        open(os.path.join(ROOT, "study", "test_unseal_log.jsonl"), "a").write(json.dumps(rec) + "\n")
        print(f"TEST SET UNSEALED (logged): {rec}\n")

    cells, _, _ = an.load_runs("jobs/p1-*")
    t = {"built": datetime.datetime.now().isoformat(timespec="seconds"),
         "sources": {"table_a": "jobs/p1-* (GPT-5.6-Luna), "
                                + ("development and test tasks" if a.include_test else "development tasks only"),
                     "table_b": os.path.basename(QWEN_XLSX)},
         "includes_test": a.include_test,
         "table_a": table_a(tasks, split_of, cells), "table_b": table_b()}
    if not a.include_test:
        leaked = set(r["task"] for r in t["table_a"]) & set(split["test"])
        if leaked:
            raise SystemExit(f"test tasks in table A: {leaked}")
    t["table_sha256"] = hashlib.sha256(json.dumps([t["table_a"], t["table_b"]], sort_keys=True).encode()).hexdigest()[:16]
    t["frozen"] = True
    json.dump(t, open(OUT_JSON, "w"), indent=1)
    open(OUT_MD, "w").write(render_md(t))
    import table_router
    table_router.write_system_prompt()
    best = {h: sum(r["best"] == h for r in t["table_a"]) for h in HARNESSES}
    by_split = {s: sum(r["split"] == s for r in t["table_a"]) for s in ("dev", "test")}
    print(f"table A: {len(t['table_a'])} tasks ({by_split['dev']} dev, {by_split['test']} test), "
          f"best harness counts {best}")
    print(f"table B: {len(t['table_b'])} Qwen tasks")
    print(f"frozen sha256 {t['table_sha256']} -> {os.path.relpath(OUT_JSON, ROOT)}, {os.path.relpath(OUT_MD, ROOT)}")


if __name__ == "__main__":
    main()
