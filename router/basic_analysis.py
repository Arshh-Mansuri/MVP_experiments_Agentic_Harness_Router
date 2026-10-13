"""Every number in router/basic_system_prompt.txt, recomputed from the 89-task benchmark workbook only
(Qwen_89_Task_Comparison_2026-09-20.xlsx, sheet "Trial details": one Qwen3-Coder trial per task and harness).
Luna's own runs (the 45 Phase 1 tasks, jobs/p1-*) are deliberately not used.

PASS/FAIL/ERROR is the workbook's status column, as on its Results sheet (pi 16 PASS; one pi trial has passed=1 with a
non-PASS status, the "pi scoring conflict" the workbook keeps). Category and difficulty come from the workbook; the
agent time limit and the instruction keywords come from each task's own task.toml and instruction.md.

  python3 router/basic_analysis.py      no network, no API calls
"""
import collections, glob, os, re, statistics as st, sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import basic_router
import profile_router as pr

WORKBOOK = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "Qwen_89_Task_Comparison_2026-09-20.xlsx")
H = ["terminus-2", "mini-swe-agent", "pi"]
KEYWORDS = {"runs a server, daemon or service": r"\b(server|daemon|service|port \d+|listen|nginx|http)",
            "virtual machine / QEMU": r"\b(qemu|virtual machine|vm|boot)\b",
            "images, video or PDFs": r"\.(png|jpe?g|gif|bmp|mp4|pdf)\b|\bimage\b|\bvideo\b",
            "compiles or builds code": r"\b(compile|build|make|cmake|gcc)\b",
            "trains or loads ML models": r"\b(train|pytorch|torch|model|huggingface)\b",
            "git": r"\bgit\b"}


def load():
    rows = [r for r in load_workbook(WORKBOOK, read_only=True, data_only=True)["Trial details"].iter_rows(values_only=True)
            if any(r)]
    i = next(k for k, r in enumerate(rows) if r[0] == "task")
    recs = [dict(zip(rows[i], r)) for r in rows[i + 1:] if r[0] and r[1] in H]
    tasks = collections.defaultdict(dict)
    for r in recs:
        tasks[r["task"]][r["harness"]] = r
    return tasks


def counts(name, tasks, keys):
    n = lambda h, s: sum(tasks[t][h]["status"] == s for t in keys)
    print(f"  {name:36} n={len(keys):2}  pass " + ", ".join(f"{h} {n(h, 'PASS')}" for h in H)
          + "   error " + ", ".join(f"{h} {n(h, 'ERROR')}" for h in H))


def main():
    tasks = load()
    allt = sorted(tasks)
    print(f"{len(allt)} tasks x {len(H)} harnesses from {os.path.basename(WORKBOOK)}\n")
    counts("overall", tasks, allt)
    for h in H:
        rs = [tasks[t][h] for t in allt]
        share = st.mean((r["tests_passed"] or 0) / r["tests_total"] for r in rs if r["tests_total"])
        cost = st.median(r["cost_usd"] for r in rs if isinstance(r["cost_usd"], (int, float)))
        print(f"  {h:15} mean share of tests passed {share:.2f}, median cost ${cost:.3f}")
    for key in ("difficulty", "category"):
        print(f"\nby {key}")
        groups = collections.defaultdict(list)
        for t in allt:
            groups[tasks[t]["mini-swe-agent"][key]].append(t)
        for k, ts in sorted(groups.items(), key=lambda x: -len(x[1])):
            counts(str(k), tasks, ts)
    print("\nby agent time limit (task.toml)")
    minutes = {t: basic_router.task_meta(t)["agent_minutes"] or 0 for t in allt}
    for lo, hi, label in ((0, 15, "up to 15 min"), (16, 30, "15 to 30 min"), (31, 10 ** 6, "over 30 min")):
        counts(label, tasks, [t for t in allt if lo <= minutes[t] <= hi])
    print("\nby instruction keywords")
    text = {t: pr.task_text(t).lower() for t in allt}
    for label, rx in KEYWORDS.items():
        counts(label, tasks, [t for t in allt if re.search(rx, text[t])])
    print("\nsetup failures (error summary 'Setup exit')")
    for t in allt:
        bad = [h for h in H if "Setup exit" in str(tasks[t][h]["error_summary"] or "")]
        if bad:
            print(f"  {t}: {', '.join(bad)} could not install")


if __name__ == "__main__":
    main()
