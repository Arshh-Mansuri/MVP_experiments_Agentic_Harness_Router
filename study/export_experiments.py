"""Export all iLab experiment data to one Excel workbook plus a CSV of the Phase 1 runs.

Everything is visible. The Phase 1 test rewards used to be hidden ("sealed") so the 18 test tasks could measure
generalisation, but they were deliberately unsealed to hard-code the best harness for all 45 tasks in the
deployment lookup (study/test_unseal_log.jsonl records every unsealing). All 45 are training data now, so there
is nothing left to seal; the 44 Phase 2 tasks are the held-out set.

Usage:
  python3 study/export_experiments.py
"""
import collections
import csv
import datetime
import glob
import json
import os
import re
import statistics as st

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HARNESSES = ["terminus-2", "mini-swe-agent", "pi"]
MODEL = "openai/gpt-5.6-luna"
SET_ASIDE_DIR = "jobs_set_aside"
OUT_XLSX = os.path.join(ROOT, "study", "ilab_experiments.xlsx")
OUT_CSV = os.path.join(ROOT, "study", "ilab_phase1_runs.csv")
PASS_GAIN = 5.0


def load_json(path):
    with open(path) as f:
        return json.load(f)


def minutes(block):
    block = block or {}
    a, b = block.get("started_at"), block.get("finished_at")
    if not a or not b:
        return None
    t0 = datetime.datetime.fromisoformat(a.replace("Z", "+00:00"))
    t1 = datetime.datetime.fromisoformat(b.replace("Z", "+00:00"))
    return round((t1 - t0).total_seconds() / 60, 2)


def parse_p1_job(job):
    """p1-<harness>-<task>-r<rep>[-retryN] -> (harness, task, rep, attempt)."""
    m = re.match(r"^p1-(.+)-r(\d+)(?:-retry(\d+))?$", job)
    if not m:
        return None
    rest, rep, retry = m.group(1), int(m.group(2)), int(m.group(3) or 0)
    for h in HARNESSES:
        if rest.startswith(h + "-"):
            return h, rest[len(h) + 1:], rep, retry
    return None


def run_row(result_path):
    r = load_json(result_path)
    a = r.get("agent_info") or {}
    ar = r.get("agent_result") or {}
    exc = r.get("exception_info") or {}
    reward = ((r.get("verifier_result") or {}).get("rewards") or {}).get("reward")
    msg_lines = (exc.get("exception_message") or "").strip().splitlines()
    return {
        "trial": r.get("trial_name"),
        "task": (r.get("task_id") or {}).get("name") or (r.get("trial_name") or "").rsplit("__", 1)[0],
        "agent": a.get("name"),
        "agent_version": a.get("version"),
        "model": (a.get("model_info") or {}).get("name"),
        "reward": reward,
        "error_type": exc.get("exception_type"),
        "error_message": msg_lines[0][:300] if msg_lines else None,
        "input_tokens": ar.get("n_input_tokens"),
        "cached_tokens": ar.get("n_cache_tokens"),
        "output_tokens": ar.get("n_output_tokens"),
        "cost_usd": ar.get("cost_usd"),
        "started_at": r.get("started_at"),
        "finished_at": r.get("finished_at"),
        "total_min": minutes(r),
        "env_setup_min": minutes(r.get("environment_setup")),
        "agent_setup_min": minutes(r.get("agent_setup")),
        "agent_run_min": minutes(r.get("agent_execution")),
        "verifier_min": minutes(r.get("verifier")),
    }


def collect_phase1(split):
    dev, test = set(split["dev"]), set(split["test"])
    paths = [(p, "") for p in glob.glob(f"{ROOT}/jobs/p1-*/*__*/result.json")]
    paths += [(p, os.path.basename(os.path.dirname(os.path.dirname(os.path.dirname(p)))))
              for p in glob.glob(f"{ROOT}/{SET_ASIDE_DIR}/*/p1-*/*__*/result.json")]
    rows = []
    for path, aside in paths:
        job = os.path.basename(os.path.dirname(os.path.dirname(path)))
        parsed = parse_p1_job(job)
        if not parsed:
            continue
        harness, task, rep, attempt = parsed
        row = run_row(path)
        split_name = "dev" if task in dev else "test" if task in test else "other"
        counted, why = True, ""
        if row["model"] != MODEL:
            counted, why = False, f"executor was {row['model']}, not Luna"
        elif aside:
            counted, why = False, f"set aside ({aside})"
        elif row["reward"] is None:
            counted, why = False, "no reward (run broke before grading)"
        elif task == "qemu-startup":
            counted, why = False, "qemu-startup dropped from scoring (pi/mini-swe-agent cannot install)"
        rows.append({
            "job": job, "split": split_name, "task": task, "harness": harness, "repeat": rep,
            "attempt": "first" if attempt == 0 else f"retry {attempt}",
            **{k: v for k, v in row.items() if k not in ("task", "agent")},
            "counted_in_scoring": "yes" if counted else "no",
            "not_counted_reason": why,
        })
    rows.sort(key=lambda x: (x["split"], x["task"], x["harness"], x["repeat"], x["attempt"]))
    return rows


def task_harness_table(rows):
    cells = collections.defaultdict(list)
    for r in rows:
        if r["counted_in_scoring"] == "yes":
            cells[(r["split"], r["task"], r["harness"])].append(r)
    out = []
    for (sp, task, h), rs in sorted(cells.items()):
        out.append({
            "split": sp, "task": task, "harness": h, "valid_runs": len(rs),
            "passes": sum(r["reward"] for r in rs),
            "pass_rate": round(st.mean(r["reward"] for r in rs), 3),
            "avg_cost_usd": round(st.mean(r["cost_usd"] or 0 for r in rs), 4),
            "avg_total_min": round(st.mean(r["total_min"] or 0 for r in rs), 1),
            "same_result_every_repeat": "yes" if len({r["reward"] for r in rs}) == 1 else "no",
        })
    return out


def strategy_rows():
    out = []
    for sp in ("dev", "test"):
        path = f"{ROOT}/study/phase1_results_{sp}.json"
        if not os.path.exists(path):
            continue
        res = load_json(path)
        base = res["results"][f"always {res['baseline']}"]["pass_rate"]
        for name, s in res["results"].items():
            out.append({
                "split": sp, "strategy": name, "baseline": res["baseline"], "tasks_scored": len(res["tasks"]),
                "pass_rate_pct": round(s["pass_rate"], 1), "passes": round(s["passes"], 2),
                "total_cost_usd": round(s["cost"], 4), "cost_per_pass_usd": round(s["cost_per_pass"], 4),
                "diff_vs_baseline_pts": round(s["diff_mean"], 1),
                "ci95_low": round(s["ci"][0], 1), "ci95_high": round(s["ci"][1], 1),
                "cost_per_pass_change_pct": round(100 * s["cost_per_pass_change"], 1),
                "meets_pass_criterion": s["meets_pass_criterion"], "meets_cost_criterion": s["meets_cost_criterion"],
            })
        if sp == "dev" and "perfect picking" in res["results"]:
            gap = res["results"]["perfect picking"]["pass_rate"] - base
            out.append({"split": sp, "strategy": "GO / NO-GO",
                        "baseline": res["baseline"],
                        "diff_vs_baseline_pts": round(gap, 1),
                        "meets_pass_criterion": "GO" if gap >= PASS_GAIN else "NO-GO"})
    return out


def kfold_rows():
    """Cross-validated scores from study/kfold_eval.py, the main evaluation. 'fixed' is the baseline used for claims."""
    out = []
    for pool in ("all", "dev"):
        for mode in ("fixed", "fold"):
            path = f"{ROOT}/study/kfold_results_{pool}_{mode}.json"
            if not os.path.exists(path):
                continue
            res = load_json(path)
            for name, s in res["results"].items():
                out.append({
                    "task_pool": pool, "baseline_mode": mode, "strategy": name,
                    "folds": res["folds"], "seed": res["seed"],
                    "tasks_scored": len(res["tasks"]), "pass_rate_pct": round(s["pass_rate"], 1),
                    "diff_vs_baseline_pts": round(s["diff_mean"], 1),
                    "ci95_low": round(s["ci"][0], 1), "ci95_high": round(s["ci"][1], 1),
                    "headroom_recovered_pct": round(100 * s["headroom_recovered"], 0),
                    "total_cost_usd": round(s["cost"], 4), "cost_per_pass_usd": round(s["cost_per_pass"], 4),
                    "cost_per_pass_change_pct": round(100 * s["cost_per_pass_change"], 1),
                    "routing_cost_usd": round(s["routing_cost"], 4),
                    "meets_pass_criterion": s["meets_pass_criterion"],
                    "meets_cost_criterion": s["meets_cost_criterion"],
                })
    return out


def qwen_rows():
    """study/qwen_vs_luna.py: does the hard-coded lookup reach the oracle, and does Qwen's run predict Luna's?"""
    path = f"{ROOT}/study/qwen_vs_luna_results.json"
    if not os.path.exists(path):
        return []
    d = load_json(path)
    out = []
    labels = [("stable_oracle_pass_rate", "stable oracle (best of 3 harnesses)"),
              ("lookup45_pass_rate", "lookup, all 45 tasks hard-coded"),
              ("lookup_dev_only_pass_rate", "lookup, development tasks only (previous)"),
              ("always_mini_pass_rate", "always mini-swe-agent")]
    for scope, v in d["verification"].items():
        for key, label in labels:
            out.append({"section": "1. lookup vs oracle", "scope": scope, "item": label,
                        "tasks": v["tasks"], "pass_rate_pct": round(v[key], 1),
                        "note": "equals the oracle by construction" if key == "lookup45_pass_rate" else ""})
    for h, verdicts in d["agreement"].items():
        for verdict in ("pass", "fail", "no signal"):
            v = verdicts[verdict]
            out.append({"section": "2. does Qwen predict Luna", "scope": h,
                        "item": f"Qwen {verdict.upper()}" if verdict != "no signal" else "Qwen gave no signal",
                        "tasks": v["tasks"], "pass_rate_pct": round(v["luna_pass_rate"], 1),
                        "note": "Luna pass rate on those tasks"})
    hit = d["best_harness_hit"]
    for key, label in (("any_passer", "Qwen passed with any harness"),
                       ("one_passer", "Qwen passed with exactly one harness")):
        n, total = hit[key]
        out.append({"section": "2. does Qwen predict Luna", "scope": "best-harness hit rate", "item": label,
                    "tasks": total, "pass_rate_pct": round(100 * n / total, 1),
                    "note": f"{n}/{total} are Luna's best harness; 33.3% by chance"})
    out.append({"section": "2. does Qwen predict Luna", "scope": "pooled over harness-task pairs",
                "item": "Luna pass rate, Qwen PASS minus Qwen FAIL", "pass_rate_pct": round(d["pooled_lift_pts"], 1),
                "note": "points; Qwen carries real signal about Luna"})
    for mode, m in d["modes"].items():
        for scope, rules in (("all 45 tasks", m["rules_all_tasks"]),
                             (f"{d['qwen_acted_tasks']} tasks Qwen could act on", m["rules_qwen_acted"])):
            for name, s in rules.items():
                if mode == "loto" and len(m["default_counts"]) > 1:
                    note = "unstable reference, sensitivity check only"
                elif name == "first Qwen pass in a fixed order (current)":
                    note = "the rule lookup() uses (by decision; CI includes 0)"
                else:
                    note = "best by the CI rule" if name == d["winner"] else ""
                out.append({"section": "3. fallback rules", "scope": f"{mode} default, {scope}", "item": name,
                            "tasks": s["tasks"], "pass_rate_pct": round(s["pass_rate"], 1),
                            "diff_vs_reference_pts": round(s["diff_vs_reference_pts"], 1),
                            "ci95_low": round(s["ci95"][0], 1), "ci95_high": round(s["ci95"][1], 1),
                            "cost_per_pass_usd": round(s["cost_per_pass"], 4),
                            "tasks_diverted_from_default": s["diverts"], "note": note})
    return out


def fallback_rows():
    """Fallback simulation from study/fallback_sim.py, with the same-harness-retry control beside it."""
    path = f"{ROOT}/study/fallback_results_dev.json"
    if not os.path.exists(path):
        return []
    res = load_json(path)
    out = []
    for name, r in res["results"].items():
        off, on, ctl = r["no_fallback"], r["fallback"], r["same_harness_retry"]
        out.append({
            "start_harness": name, "tasks_scored": len(res["tasks"]),
            "retry_order": " -> ".join(res["retry_order"]), "max_retries": res["max_retries"],
            "pass_rate_no_fallback_pct": round(off["pass_rate"], 1),
            "pass_rate_with_fallback_pct": round(on["pass_rate"], 1),
            "gain_switch_harness_pts": round(on["pass_rate"] - off["pass_rate"], 1),
            "gain_same_harness_retry_pts": round(ctl["pass_rate"] - off["pass_rate"], 1),
            "triggers": ", ".join(f"{k} {v}" for k, v in sorted(r["triggers"].items())) or "none",
            "retries_used": round(on["retries"], 2), "passes_gained": round(r["passes_gained"], 2),
            "extra_cost_usd": round(on["cost"] - off["cost"], 4),
            "extra_cost_per_gained_pass_usd": (round(r["extra_cost_per_gained_pass"], 4)
                                               if r["extra_cost_per_gained_pass"] else None),
            "extra_hours_per_gained_pass": (round(r["extra_hours_per_gained_pass"], 2)
                                            if r["extra_hours_per_gained_pass"] else None),
        })
    return out


def feature_rows():
    path = f"{ROOT}/study/task_features.csv"
    if not os.path.exists(path):
        return []
    with open(path) as f:
        return list(csv.DictReader(f))


def pick_rows():
    path = f"{ROOT}/router/picks_profiles.json"
    if not os.path.exists(path):
        return []
    rows = []
    for r in load_json(path):
        tok = r.get("tokens")
        rows.append({
            "router": r.get("router"), "task": r.get("task"), "repeat": r.get("rep"),
            "picked_harness": r.get("harness"), "underlying_model": r.get("underlying_model"),
            "reason": r.get("reason"),
            "tokens": json.dumps(tok) if isinstance(tok, dict) else tok,
            "cost_usd": r.get("cost"), "seconds": r.get("seconds"),
            "profiles_sha256": r.get("profiles_sha256"),
        })
    return sorted(rows, key=lambda x: (x["router"] or "", x["task"] or "", x["repeat"] or 0))


def failure_rows(rows):
    counts = collections.Counter()
    attempts = collections.Counter()
    for r in rows:
        attempts[r["harness"]] += 1
        if r["reward"] is None:
            counts[(r["harness"], r["error_type"] or "no reward")] += 1
    out = [{"harness": h, "error_type": e, "attempts_failed": n,
            "share_of_harness_attempts_pct": round(100 * n / attempts[h], 1)}
           for (h, e), n in sorted(counts.items())]
    for h in HARNESSES:
        broke = sum(n for (hh, _), n in counts.items() if hh == h)
        out.append({"harness": h, "error_type": "TOTAL (any error)", "attempts_failed": broke,
                    "share_of_harness_attempts_pct": round(100 * broke / attempts[h], 1) if attempts[h] else None})
    return out


def older_rows():
    rows = []
    for path in glob.glob(f"{ROOT}/jobs/*/*__*/result.json"):
        job = os.path.basename(os.path.dirname(os.path.dirname(path)))
        if job.startswith("p1-"):
            continue
        try:
            row = run_row(path)
        except (json.JSONDecodeError, OSError):
            continue
        rows.append({"job": job, "status": "pre-study (not part of Phase 1)", **row})
    return sorted(rows, key=lambda x: (x["started_at"] or "", x["job"]))


def notes_rows(split, rows):
    prof_path = f"{ROOT}/router/harness_profiles.json"
    prof = load_json(prof_path) if os.path.exists(prof_path) else {}
    table_path = f"{ROOT}/router/success_table.json"
    table = load_json(table_path) if os.path.exists(table_path) else {}
    luna = [r for r in rows if r["model"] == MODEL]
    return [
        ("Generated", datetime.datetime.now().isoformat(timespec="seconds")),
        ("Test set", f"UNSEALED. The {len(split['test'])} test tasks were held back to measure generalisation, then "
                     "deliberately unsealed so the deployment lookup could hard-code the best harness for all "
                     f"{len(split['dev']) + len(split['test'])} tasks. All of them are training data now and none of "
                     "them can measure generalisation any more; the 44 Phase 2 tasks are the held-out set. "
                     "study/test_unseal_log.jsonl records every unsealing."),
        ("Executor model", f"openrouter/{MODEL} (fixed for every run)"),
        ("Harnesses compared", ", ".join(HARNESSES)),
        ("Design", f"{len(split['dev'])} dev + {len(split['test'])} test tasks x {len(HARNESSES)} harnesses x "
                   f"{split['reps']} repeats = {(len(split['dev']) + len(split['test'])) * len(HARNESSES) * split['reps']} slots"),
        ("Task selection seed", split.get("seed")),
        ("Phase 1 attempts in this file", f"{len(rows)} ({len(luna)} with Luna)"),
        ("Retries", "A slot that broke before grading is retried up to 2 more times; broken attempts are not scored."),
        ("Scoring", "Pass rate = mean reward over valid repeats per task, then averaged over tasks. "
                    "Baseline = best single harness on the dev set. 95% CIs from 10,000 paired bootstrap resamples over tasks."),
        ("Success criteria", "Router beats baseline by >= 5 points with CI above 0, or cuts cost per pass by >= 30% "
                             "while staying within 3 points."),
        ("Routers", "jev = Jev Router (typesafe/jev-router) picks using Luna's frozen harness profiles, kept as a record "
                    "only: that router forwards to other models and was retired on 7 Oct; luna-profiles = Luna picks "
                    "using the same profiles; luna-cards = Luna picks using the earlier hand-written harness cards."),
        ("Harness profiles", f"frozen={prof.get('frozen')}, sha256={prof.get('profiles_sha256')}"),
        ("Decision: qemu-startup", "Dropped from scoring (decided before unsealing): pi and mini-swe-agent fail their own "
                                   "install step in that task's container; only terminus-2 runs."),
        ("Decision: sleep-affected runs", "3 runs in progress while the laptop slept (5 Oct, 00:46-09:49) were set aside "
                                          "in jobs_set_aside/ and rerun. They are listed but not counted."),
        ("Withdrawn", "The earlier Jev-as-picker study (29 Sep) was withdrawn and is not in this file."),
        ("Sheets", "Phase1 runs: one row per attempt | Task x harness: per-cell pass rates | Strategies: single dev/test "
                   "split | K-fold: cross-validated scores, the main evaluation | Qwen vs Luna: lookup-vs-oracle check "
                   "and how well a different model's run predicts Luna's | Fallback: retry on an observable failure, "
                   "with the same-harness control | Task features: task.toml metadata and keyword flags | "
                   "Router picks: every pick with reason | Failures: broken attempts by harness | Older runs: pre-study jobs"),
        ("Deployment lookup", f"router/success_table.json, frozen at sha256={table.get('table_sha256')}, Table A covers "
                              f"{len(table.get('table_a') or [])} tasks. table_router.lookup() answers those from the "
                              "table with no model call and reaches the best-of-3-harnesses oracle on them by "
                              "construction (70.4% pass against 54.8% for always mini-swe-agent). That is memorisation, "
                              "not generalisation: an unlisted task gets the Qwen fallback or the default harness."),
        ("Qwen as a fallback hint", "Qwen3-Coder's 89-task run does predict Luna: Luna passes 36.3 points more often "
                                    "with a harness Qwen passed with, and when Qwen passes at all its harness is Luna's "
                                    "best one 76% of the time against 33% by chance. But as a fallback rule it only "
                                    "changes the pick on 5 of the 45 tasks and gains 2.2 points, CI [-2.2, +7.4], so "
                                    "it is unproven. lookup() uses it anyway as a judgement call: it leans "
                                    "positive, is cheaper per pass, and changes only 2 of the 44 unrun tasks. "
                                    "It also skips harnesses that could not install for the task in the Qwen run."),
        ("K-fold task pools", "task_pool='dev' is the original 27-task evaluation; 'all' adds the unsealed test tasks "
                              "(44 of 45 - qemu-startup is excluded because only terminus-2 ever ran it, so no strategy "
                              "has a real choice there)."),
        ("Read K-fold, not Strategies", "Strategies scores one dev/test split and lets a strategy that learned from those "
                                        "same tasks look perfect. K-fold refits anything learned on the training folds only. "
                                        "Use baseline_mode='fixed' rows for claims: 'fold' re-picks the best single harness "
                                        "per fold and is itself unstable, which flatters every strategy."),
        ("Fallback: read the control", "gain_same_harness_retry_pts is what simply rerunning the same harness achieves. Most "
                                       "triggers are infrastructure crashes, so a switch only earns credit where it beats "
                                       "that column. For mini-swe-agent it does not."),
        ("Harness versions", "Phase 1 mixes pi 1.0.0/1.0.1/1.0.2 (terminus-2 2.0.0 and mini-swe-agent 2.4.6 were stable), "
                             "so 'pi' is not one harness here. Future runs pin the version via run_phase1.sh."),
        ("Column: reward", "1 = task solved, 0 = not solved, blank = run broke before grading"),
        ("Column: times", "Minutes. env_setup = container start, agent_setup = harness install, agent_run = the agent "
                          "working, verifier = grading."),
    ]


HEADER_FILL = PatternFill("solid", fgColor="1F3D36")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def add_sheet(wb, title, rows, widths=None):
    ws = wb.create_sheet(title)
    if not rows:
        ws.append(["(no data yet)"])
        return ws
    cols = list(dict.fromkeys(k for r in rows for k in r))
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c) for c in cols])
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        width = (widths or {}).get(c) or min(60, max(len(c), *(len(str(r.get(c) or "")) for r in rows[:300])) + 2)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    split = load_json(f"{ROOT}/study/phase1_tasks.json")
    runs = collect_phase1(split)

    wb = Workbook()
    ws = wb.active
    ws.title = "Notes"
    for k, v in notes_rows(split, runs):
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110

    add_sheet(wb, "Phase1 runs", runs, {"error_message": 60})
    add_sheet(wb, "Task x harness", task_harness_table(runs))
    add_sheet(wb, "Strategies", strategy_rows())
    add_sheet(wb, "K-fold", kfold_rows())
    add_sheet(wb, "Qwen vs Luna", qwen_rows(), {"item": 44, "scope": 36, "note": 48})
    add_sheet(wb, "Fallback", fallback_rows(), {"triggers": 40, "retry_order": 40})
    add_sheet(wb, "Task features", feature_rows())
    add_sheet(wb, "Router picks", pick_rows(), {"reason": 80})
    add_sheet(wb, "Failures", failure_rows(runs))
    add_sheet(wb, "Older runs", older_rows(), {"error_message": 60})
    wb.save(OUT_XLSX)

    if runs:
        with open(OUT_CSV, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(dict.fromkeys(k for r in runs for k in r)))
            w.writeheader()
            w.writerows(runs)

    print(f"wrote {os.path.relpath(OUT_XLSX, ROOT)} and {os.path.relpath(OUT_CSV, ROOT)}: "
          f"{len(runs)} Phase 1 attempts, test rewards visible")


if __name__ == "__main__":
    main()
