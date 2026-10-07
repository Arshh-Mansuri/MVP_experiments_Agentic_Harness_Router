"""What else the Qwen3-Coder 89-task workbook can tell us about Luna, beyond binary PASS/FAIL.

The workbook (archive/Qwen_89_Task_Comparison_2026-09-20.xlsx, sheet "Trial details") has one Qwen trial per task and
harness with test counts, tokens, cost, time, exception type, category and difficulty. Luna's side comes from
study/ilab_phase1_runs.csv (453 attempts). Everything is local; no API key.

  1. harness leaderboard        does the harness ranking carry over from Qwen to Luna?
  2. task difficulty            does Qwen predict how hard a task is for Luna, and does partial credit
                                (tests passed / tests total) predict it better than PASS/FAIL?
  3. harness choice per task    a fallback that ranks harnesses by Qwen's partial credit, scored against
                                always mini-swe-agent like study/qwen_vs_luna.py
  4. task size and cost         does Qwen's token use predict Luna's cost? If so, forecast Phase 2
  5. harness cost premium       is terminus-2's extra cost a harness property or a Luna property?
  6. crashes                    does a Qwen ERROR on a task x harness predict a Luna crash on the same cell?
  7. difficulty label           does Terminal-Bench's own easy/medium/hard label add anything?
  8. opencode                   the fourth Qwen harness, which Luna never ran
  9. Phase 2 triage             the 44 unrun tasks ranked by what the above says

usage: python3 study/qwen_deep.py      writes study/qwen_deep_results.json
"""
import collections
import csv
import json
import os
import random
import statistics as st

from openpyxl import load_workbook
from scipy.stats import mannwhitneyu, spearmanr

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "archive", "Qwen_89_Task_Comparison_2026-09-20.xlsx")
RUNS = os.path.join(ROOT, "study", "ilab_phase1_runs.csv")
OURS = ("terminus-2", "mini-swe-agent", "pi")
ALL4 = OURS + ("opencode",)
DEFAULT = "mini-swe-agent"
LUNA = "openai/gpt-5.6-luna"
BOOT = 10000


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def load_qwen():
    ws = load_workbook(XLSX, data_only=True)["Trial details"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "task")
    cols = rows[hi]
    q = {}
    for r in rows[hi + 1:]:
        if not r[0]:
            continue
        d = dict(zip(cols, r))
        tp, tt = num(d["tests_passed"]), num(d["tests_total"])
        q[(d["task"], d["harness"])] = {
            "status": d["status"], "pass": 1.0 if d["status"] == "PASS" else 0.0, "reward": num(d["reward"]),
            "frac": (tp / tt) if tp is not None and tt else (1.0 if d["status"] == "PASS" else 0.0),
            "tokens": num(d["total_tokens"]), "cost": num(d["cost_usd"]), "secs": num(d["total_s"]),
            "exception": d["exception_type"], "error_summary": d["error_summary"],
            "category": d["category"], "difficulty": d["difficulty"]}
    return q


def load_luna():
    """(task, harness) -> scored rewards, costs, tokens, minutes; plus crash counts over all non-set-aside attempts."""
    cells = collections.defaultdict(lambda: {"rewards": [], "cost": [], "tokens": [], "min": [], "attempts": 0,
                                             "crashes": 0})
    for r in csv.DictReader(open(RUNS)):
        if r["model"] != LUNA or r["not_counted_reason"].startswith("set aside"):
            continue
        c = cells[(r["task"], r["harness"])]
        c["attempts"] += 1
        if r["reward"] == "":
            c["crashes"] += 1
        if r["counted_in_scoring"] == "yes":
            c["rewards"].append(float(r["reward"]))
            c["cost"].append(num(r["cost_usd"]) or 0.0)
            c["tokens"].append((num(r["input_tokens"]) or 0) + (num(r["output_tokens"]) or 0))
            c["min"].append(num(r["total_min"]) or 0.0)
    return cells


def mean(xs):
    xs = [x for x in xs if x is not None]
    return st.mean(xs) if xs else None


def auc(pos, neg):
    """Probability a random positive outranks a random negative (0.5 = no signal)."""
    if not pos or not neg:
        return None
    return mannwhitneyu(pos, neg).statistic / (len(pos) * len(neg))


def boot_diff(a, b, seed=0):
    rnd = random.Random(seed)
    d = [x - y for x, y in zip(a, b)]
    ms = sorted(st.mean(rnd.choice(d) for _ in d) for _ in range(BOOT))
    return st.mean(d), (ms[int(0.025 * BOOT)], ms[int(0.975 * BOOT)])


def rho(xs, ys):
    pairs = [(x, y) for x, y in zip(xs, ys) if x is not None and y is not None]
    r, p = spearmanr([x for x, _ in pairs], [y for _, y in pairs])
    return {"rho": round(float(r), 2), "p": round(float(p), 4), "n": len(pairs)}


def main():
    q, L = load_qwen(), load_luna()
    qtasks = sorted({t for t, _ in q})
    luna_tasks = sorted({t for (t, h), c in L.items() if c["attempts"]})
    both = [t for t in luna_tasks if t in qtasks and any(L[(t, h)]["rewards"] for h in OURS)]
    unrun = [t for t in qtasks if t not in luna_tasks]
    n = len(both)
    lp = lambda t, h: mean(L[(t, h)]["rewards"]) if L[(t, h)]["rewards"] else None
    luna_task = {t: mean([lp(t, h) for h in OURS]) for t in both}
    out = {"qwen_tasks": len(qtasks), "luna_tasks": len(luna_tasks), "overlap": len(both), "unrun": len(unrun)}
    print(f"Qwen: {len(qtasks)} tasks x {len(ALL4)} harnesses, one trial each. Luna ran {len(luna_tasks)} tasks, "
          f"{n} scored (qemu-startup is excluded from scoring). Unrun by Luna: {len(unrun)}.")

    # 1. leaderboard
    print("\n=== 1. Harness leaderboard: does the ranking carry over? ===")
    print(f"  {'harness':16}{'Qwen, 89':>10}{f'Qwen, {n}':>10}{f'Luna, {n}':>10}")
    board = {}
    for h in ALL4:
        q89 = 100 * mean([q[(t, h)]["pass"] for t in qtasks])
        q45 = 100 * mean([q[(t, h)]["pass"] for t in both])
        l45 = 100 * mean([lp(t, h) or 0 for t in both]) if h in OURS else None
        board[h] = {"qwen_89": q89, "qwen_45": q45, "luna_45": l45}
        print(f"  {h:16}{q89:9.1f}%{q45:9.1f}%" + (f"{l45:9.1f}%" if l45 is not None else "   (never run)"))
    out["leaderboard"] = board

    # 2. difficulty
    print("\n=== 2. Task difficulty: does Qwen predict how hard a task is for Luna? ===")
    signals = {
        "PASS count over our 3 harnesses": {t: sum(q[(t, h)]["pass"] for h in OURS) for t in both},
        "PASS count over all 4 harnesses": {t: sum(q[(t, h)]["pass"] for h in ALL4) for t in both},
        "partial credit, mean over 4 harnesses": {t: mean([q[(t, h)]["frac"] for h in ALL4]) for t in both},
        "partial credit, best of 4 harnesses": {t: max(q[(t, h)]["frac"] for h in ALL4) for t in both},
        "tokens used (fewer = easier), median of 4": {t: -(st.median([q[(t, h)]["tokens"] or 0 for h in ALL4]))
                                                     for t in both},
    }
    solved = {t: luna_task[t] >= 0.5 for t in both}
    diff = {}
    print(f"  {'Qwen signal':44}{'Spearman vs Luna pass':>24}{'AUC: Luna solves it':>22}")
    for name, s in signals.items():
        r = rho([s[t] for t in both], [luna_task[t] for t in both])
        a = auc([s[t] for t in both if solved[t]], [s[t] for t in both if not solved[t]])
        diff[name] = {**r, "auc": a}
        print(f"  {name:44}{r['rho']:>+10.2f} (p={r['p']:.3f}){a:>18.2f}")
    print(f"  (AUC 0.5 = no signal, 1.0 = perfect; Luna 'solves' = mean pass >= 0.5 over its 3 harnesses, "
          f"{sum(solved.values())}/{len(both)} tasks)")
    out["difficulty"] = diff

    # 3. harness choice from partial credit
    print("\n=== 3. Picking the harness from Qwen's partial credit (vs always mini-swe-agent) ===")
    def pick_binary(t):
        passed = [h for h in (DEFAULT, "terminus-2", "pi") if q[(t, h)]["pass"]]
        return passed[0] if passed else DEFAULT
    def pick_frac(t):
        return max((DEFAULT, "terminus-2", "pi"), key=lambda h: q[(t, h)]["frac"])  # max keeps the first on ties
    def pick_frac_strict(t):
        best = pick_frac(t)
        return best if q[(t, best)]["frac"] > q[(t, DEFAULT)]["frac"] + 0.25 else DEFAULT
    rules = {"always mini-swe-agent": lambda t: DEFAULT, "first Qwen PASS (binary)": pick_binary,
             "highest Qwen partial credit": pick_frac,
             "highest partial credit, only if 25+ points above mini": pick_frac_strict}
    scored = [t for t in both if all(lp(t, h) is not None for h in OURS)]
    base = [lp(t, DEFAULT) for t in scored]
    pick3 = {}
    print(f"  {len(scored)} tasks where Luna ran all three harnesses")
    print(f"  {'rule':52}{'Luna pass':>10}{'diff (95% CI)':>24}{'changed':>9}")
    for name, f in rules.items():
        got = [lp(t, f(t)) for t in scored]
        d, ci = boot_diff(got, base)
        ch = sum(f(t) != DEFAULT for t in scored)
        pick3[name] = {"pass": 100 * st.mean(got), "diff": 100 * d, "ci": [100 * ci[0], 100 * ci[1]], "changed": ch}
        print(f"  {name:52}{100 * st.mean(got):9.1f}%{100 * d:+8.1f} [{100 * ci[0]:+5.1f}, {100 * ci[1]:+5.1f}]{ch:>9}")
    out["harness_choice"] = pick3

    # 4. size and cost
    print("\n=== 4. Task size: does Qwen's token use predict Luna's cost and time? ===")
    qtok = {t: st.median([q[(t, h)]["tokens"] or 0 for h in OURS]) for t in qtasks}
    qsec = {t: st.median([q[(t, h)]["secs"] or 0 for h in OURS]) for t in qtasks}
    lcost = {t: mean([mean(L[(t, h)]["cost"]) for h in OURS if L[(t, h)]["cost"]]) for t in both}
    ltok = {t: mean([mean(L[(t, h)]["tokens"]) for h in OURS if L[(t, h)]["tokens"]]) for t in both}
    lmin = {t: mean([mean(L[(t, h)]["min"]) for h in OURS if L[(t, h)]["min"]]) for t in both}
    size = {"Qwen tokens vs Luna tokens": rho([qtok[t] for t in both], [ltok[t] for t in both]),
            "Qwen tokens vs Luna cost": rho([qtok[t] for t in both], [lcost[t] for t in both]),
            "Qwen seconds vs Luna minutes": rho([qsec[t] for t in both], [lmin[t] for t in both])}
    for k, v in size.items():
        print(f"  {k:32} Spearman {v['rho']:+.2f} (p={v['p']:.4f}, n={v['n']})")
    # Ratio of totals, not a median of per-task ratios: totals are what a budget needs, and the per-task median
    # overshot Phase 1 by 27% because cheap tasks have noisy ratios.
    ratio = sum(lcost[t] for t in both) / sum(qtok[t] for t in both)
    fc = {t: ratio * qtok[t] for t in unrun}
    runs = len(unrun) * len(OURS) * 3
    est = 3 * len(OURS) * sum(fc.values())
    actual = 3 * len(OURS) * sum(lcost[t] for t in both)
    half = sorted(both)
    a, b = half[::2], half[1::2]
    r_a = sum(lcost[t] for t in a) / sum(qtok[t] for t in a)
    pred_b, act_b = r_a * sum(qtok[t] for t in b), sum(lcost[t] for t in b)
    print(f"  Phase 2 forecast ({len(unrun)} tasks x 3 harnesses x 3 repeats = {runs} runs): about ${est:.2f} of Luna "
          f"spend, against ${actual:.2f} for the {n} Phase 1 tasks (scored runs only; crashed retries add more)")
    print(f"  split-half check: fitted on {len(a)} tasks, it predicts the other {len(b)} at ${pred_b:.2f} per repeat "
          f"set against ${act_b:.2f} actual ({100 * (pred_b / act_b - 1):+.0f}%)")
    out["size"] = {**size, "phase2_forecast_usd": est, "phase1_actual_usd": actual,
                   "split_half_error_pct": 100 * (pred_b / act_b - 1)}

    # 5. cost premium
    print("\n=== 5. Harness cost premium: harness property or model property? ===")
    def premium(get, tasks):
        rs = [get(t, "terminus-2") / get(t, DEFAULT) for t in tasks if get(t, "terminus-2") and get(t, DEFAULT)]
        return st.median(rs), len(rs)
    qc = lambda t, h: q[(t, h)]["cost"]
    lc = lambda t, h: mean(L[(t, h)]["cost"]) if L[(t, h)]["cost"] else None
    prem = {}
    for label, get, tasks in (("Qwen, 89 tasks", qc, qtasks), (f"Qwen, {n} tasks", qc, both), (f"Luna, {n} tasks", lc, both)):
        m, n = premium(get, tasks)
        prem[label] = m
        print(f"  {label:16} terminus-2 costs {m:.2f}x mini-swe-agent per task (median, n={n})")
    for h in ("pi", "opencode"):
        rs = [q[(t, h)]["cost"] / q[(t, DEFAULT)]["cost"] for t in qtasks if q[(t, h)]["cost"] and q[(t, DEFAULT)]["cost"]]
        prem[f"Qwen {h}"] = st.median(rs)
        print(f"  Qwen, 89 tasks   {h} costs {st.median(rs):.2f}x mini-swe-agent per task (median, n={len(rs)})")
    tot = lambda h: sum(lc(t, h) or 0 for t in both)
    prem["Luna total terminus-2 / mini"] = tot("terminus-2") / tot(DEFAULT)
    print(f"  Luna, summed over {n} tasks: terminus-2 costs {prem['Luna total terminus-2 / mini']:.2f}x mini-swe-agent "
          f"- far above the per-task median, so the premium is concentrated in a few expensive tasks")
    out["cost_premium"] = prem

    # 6. crashes
    print("\n=== 6. Crashes: does a Qwen ERROR predict a Luna crash on the same task x harness? ===")
    exc = collections.Counter((h, q[(t, h)]["exception"]) for t in qtasks for h in ALL4 if q[(t, h)]["status"] == "ERROR")
    for (h, e), n in sorted(exc.items()):
        print(f"  Qwen ERROR  {h:15} {e or '(no exception recorded)':28} {n}")
    pairs = [(t, h) for t in both for h in OURS if L[(t, h)]["attempts"]]
    qerr = [p for p in pairs if q[p]["status"] == "ERROR"]
    rate = lambda ps: (sum(L[p]["crashes"] for p in ps) / sum(L[p]["attempts"] for p in ps)) if ps else None
    rest = [p for p in pairs if p not in qerr]
    print(f"  Luna crash rate on cells where Qwen ERRORed: {100 * rate(qerr):.0f}% of attempts ({len(qerr)} cells); "
          f"elsewhere {100 * rate(rest):.0f}% ({len(rest)} cells)")
    for p in qerr:
        print(f"    {p[0]:32} {p[1]:15} Qwen {q[p]['exception']}, Luna {L[p]['crashes']}/{L[p]['attempts']} attempts crashed")
    out["crashes"] = {"qwen_error_cells": len(qerr), "luna_crash_rate_there": rate(qerr), "luna_crash_rate_elsewhere": rate(rest)}

    # 7. difficulty label
    print("\n=== 7. Terminal-Bench's own difficulty label ===")
    lab = {t: q[(t, DEFAULT)]["difficulty"] for t in qtasks}
    labels = {}
    for d in ("easy", "medium", "hard"):
        qs = [t for t in qtasks if lab[t] == d]
        ls = [t for t in both if lab[t] == d]
        qr = 100 * mean([q[(t, h)]["pass"] for t in qs for h in OURS]) if qs else None
        lr = 100 * mean([luna_task[t] for t in ls]) if ls else None
        labels[d] = {"qwen_tasks": len(qs), "qwen_pass": qr, "luna_tasks": len(ls), "luna_pass": lr}
        print(f"  {d:7} Qwen {qr if qr is not None else float('nan'):5.1f}% over {len(qs):2} tasks | "
              f"Luna {lr if lr is not None else float('nan'):5.1f}% over {len(ls):2} tasks")
    rank = {"easy": 0, "medium": 1, "hard": 2}
    r_lab = rho([-rank[lab[t]] for t in both], [luna_task[t] for t in both])
    print(f"  label vs Luna pass: Spearman {r_lab['rho']:+.2f} (p={r_lab['p']:.3f}) - compare with section 2")
    out["difficulty_label"] = {**labels, "spearman_vs_luna": r_lab}

    # 8. opencode
    print("\n=== 8. opencode, the harness Luna never ran ===")
    only_oc = [t for t in qtasks if q[(t, "opencode")]["pass"] and not any(q[(t, h)]["pass"] for h in OURS)]
    oc_cost = mean([q[(t, "opencode")]["cost"] for t in qtasks])
    print(f"  Qwen pass {board['opencode']['qwen_89']:.1f}% on 89 tasks; mean cost ${oc_cost:.3f} per task")
    print(f"  tasks only opencode solved: {len(only_oc)} {only_oc}")
    out["opencode"] = {"only_opencode_solved": only_oc, "mean_cost": oc_cost}

    # 9. Phase 2 triage
    print(f"\n=== 9. Phase 2 triage: the {len(unrun)} tasks Luna has not run ===")
    tri = []
    for t in unrun:
        n = sum(q[(t, h)]["pass"] for h in OURS)
        best_frac = max(q[(t, h)]["frac"] for h in ALL4)
        spread = max(q[(t, h)]["frac"] for h in OURS) - min(q[(t, h)]["frac"] for h in OURS)
        group = ("routing likely matters" if n in (1, 2) or (n == 0 and spread >= 0.5) else
                 "easy: any harness, pick the cheapest" if n == 3 else
                 "near miss: Qwen got most tests" if best_frac >= 0.5 else "hard: Qwen got little")
        tri.append({"task": t, "group": group, "qwen_passes": n, "qwen_best_partial": round(best_frac, 2),
                    "partial_spread": round(spread, 2), "forecast_luna_cost_per_run": round(fc[t], 4),
                    "difficulty": lab[t], "category": q[(t, DEFAULT)]["category"]})
    order = ["routing likely matters", "near miss: Qwen got most tests", "easy: any harness, pick the cheapest",
             "hard: Qwen got little"]
    tri.sort(key=lambda r: (order.index(r["group"]), -r["partial_spread"]))
    for g in order:
        ts = [r for r in tri if r["group"] == g]
        cost = sum(r["forecast_luna_cost_per_run"] for r in ts) * 9
        print(f"  {g:38} {len(ts):2} tasks, ~${cost:.2f} to run 3x3: "
              + ", ".join(r["task"] for r in ts[:8]) + (" ..." if len(ts) > 8 else ""))
    out["phase2_triage"] = tri

    path = os.path.join(ROOT, "study", "qwen_deep_results.json")
    json.dump(out, open(path, "w"), indent=1, default=float)
    print(f"\nwrote {os.path.relpath(path, ROOT)}")


if __name__ == "__main__":
    main()
