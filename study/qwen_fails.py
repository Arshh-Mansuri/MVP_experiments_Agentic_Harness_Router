"""What the Qwen3-Coder workbook's failures say: how close they came, what they cost, how they ended, whether Luna
fails the same way, and which crashes are infrastructure rather than the agent.

Reuses the loaders in study/qwen_deep.py. Local files only.

usage: python3 study/qwen_fails.py      writes study/qwen_fails_results.json
"""
import collections
import json
import os
import statistics as st

from openpyxl import load_workbook

import qwen_deep as qd

ROOT = qd.ROOT
OURS, ALL4 = qd.OURS, qd.ALL4


def raw_rows():
    ws = load_workbook(qd.XLSX, data_only=True)["Trial details"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "task")
    return [dict(zip(rows[hi], r)) for r in rows[hi + 1:] if r[0]]


def bucket(tp, tt):
    if tt is None or tp is None:
        return "no test counts"
    if tt == 1:
        return "single test (no partial credit)"
    f = tp / tt
    return ("0 tests passed" if tp == 0 else "one test short" if tp == tt - 1 else
            "more than half passed" if f > 0.5 else "half or fewer passed")


def main():
    R = raw_rows()
    q, L = qd.load_qwen(), qd.load_luna()
    out = {}
    fails = [d for d in R if d["status"] == "FAIL"]
    errors = [d for d in R if d["status"] == "ERROR"]
    passes = [d for d in R if d["status"] == "PASS"]
    print(f"Qwen trials: {len(R)} = {len(passes)} PASS, {len(fails)} FAIL, {len(errors)} ERROR")

    # 1. how close
    print("\n=== 1. How close did the failures come? ===")
    order = ["0 tests passed", "half or fewer passed", "more than half passed", "one test short",
             "single test (no partial credit)", "no test counts"]
    by = collections.Counter(bucket(qd.num(d["tests_passed"]), qd.num(d["tests_total"])) for d in fails)
    for b in order:
        print(f"  {b:34} {by[b]:4} ({100 * by[b] / len(fails):4.1f}%)")
    print("  by harness (share of its FAILs that were one test short):")
    for h in ALL4:
        hf = [d for d in fails if d["harness"] == h]
        n1 = sum(bucket(qd.num(d["tests_passed"]), qd.num(d["tests_total"])) == "one test short" for d in hf)
        print(f"    {h:16} {len(hf):3} FAILs, {n1:2} one test short ({100 * n1 / len(hf):.0f}%)")
    out["closeness"] = dict(by)

    # 2. what failing cost
    print("\n=== 2. What did failing cost? ===")
    cost = lambda ds: sum(qd.num(d["cost_usd"]) or 0 for d in ds)
    total = cost(R)
    for label, ds in (("PASS", passes), ("FAIL", fails), ("ERROR", errors)):
        print(f"  {label:6} {len(ds):3} trials  ${cost(ds):7.2f}  ({100 * cost(ds) / total:4.1f}% of spend)  "
              f"median {st.median([qd.num(d['total_tokens']) or 0 for d in ds]):>9,.0f} tokens, "
              f"median {st.median([qd.num(d['agent_exec_s']) or 0 for d in ds]) / 60:5.1f} min")
    out["spend_share"] = {"pass": cost(passes) / total, "fail": cost(fails) / total, "error": cost(errors) / total}
    print("  per harness: share of its spend that bought a FAIL or ERROR")
    for h in ALL4:
        hs = [d for d in R if d["harness"] == h]
        bad = [d for d in hs if d["status"] != "PASS"]
        print(f"    {h:16} ${cost(hs):6.2f} total, {100 * cost(bad) / cost(hs):4.1f}% wasted")

    # 3. how failures ended: quit early or ran long, relative to a pass of the same task
    print("\n=== 3. How did failures end? (time against a harness that passed the same task) ===")
    pass_time = collections.defaultdict(list)
    for d in passes:
        pass_time[d["task"]].append(qd.num(d["agent_exec_s"]) or 0)
    shape = collections.Counter()
    examples = collections.defaultdict(list)
    for d in fails:
        if d["task"] not in pass_time:
            continue
        ref = st.median(pass_time[d["task"]])
        t = qd.num(d["agent_exec_s"]) or 0
        k = ("quit early (< half the passing time)" if t < 0.5 * ref else
                 "ran long (> 2x the passing time)" if t > 2 * ref else "similar time")
        shape[k] += 1
        examples[k].append(f"{d['task']}/{d['harness']}")
    n = sum(shape.values())
    print(f"  {n} FAILs on tasks another harness passed:")
    for k, v in shape.most_common():
        print(f"    {k:38} {v:3} ({100 * v / n:.0f}%)  e.g. {', '.join(examples[k][:3])}")
    out["fail_shape"] = dict(shape)

    # 4. does a Qwen near miss mean Luna passes the same cell?
    print("\n=== 4. Qwen FAIL cells: does Luna pass them, and does closeness matter? ===")
    cells = collections.defaultdict(list)
    for d in fails:
        h, t = d["harness"], d["task"]
        if h in OURS and L[(t, h)]["rewards"]:
            b = bucket(qd.num(d["tests_passed"]), qd.num(d["tests_total"]))
            g = ("0 tests passed" if b == "0 tests passed" else
                 "some tests passed" if b in ("half or fewer passed", "more than half passed", "one test short")
                 else "single test / no counts")
            cells[g].append(qd.mean(L[(t, h)]["rewards"]))
            if b == "one test short":
                cells["  of which one test short"].append(qd.mean(L[(t, h)]["rewards"]))
    for g in ("0 tests passed", "some tests passed", "  of which one test short", "single test / no counts"):
        xs = cells[g]
        if xs:
            print(f"  Qwen {g:28} {len(xs):3} cells -> Luna passes {100 * st.mean(xs):5.1f}% of runs")
    out["luna_on_qwen_fails"] = {g: {"cells": len(xs), "luna_pass": st.mean(xs)} for g, xs in cells.items() if xs}

    # 5. tasks Qwen failed everywhere
    print("\n=== 5. Tasks every Qwen harness failed ===")
    tasks = sorted({d["task"] for d in R})
    dead = [t for t in tasks if not any(q[(t, h)]["pass"] for h in ALL4)]
    on_luna = [t for t in dead if any(L[(t, h)]["rewards"] for h in OURS)]
    lr = {t: qd.mean([qd.mean(L[(t, h)]["rewards"]) for h in OURS if L[(t, h)]["rewards"]]) for t in on_luna}
    print(f"  {len(dead)} of {len(tasks)} tasks; Luna ran {len(on_luna)} of them and averages "
          f"{100 * st.mean(lr.values()):.0f}% on them")
    print(f"  Luna solved at least sometimes: {sum(v > 0 for v in lr.values())}/{len(on_luna)}; "
          f"Luna mostly solves (>= 2/3): {', '.join(t for t, v in sorted(lr.items()) if v >= 2 / 3)}")
    out["qwen_dead_tasks"] = {"count": len(dead), "luna_ran": len(on_luna), "luna_mean": st.mean(lr.values())}

    # 6. crashes: infrastructure or agent?
    print("\n=== 6. The 27 ERRORs: infrastructure or agent? ===")
    infra = {"Setup exit 100", "ApiRateLimitError", "ApiOverloadedError"}
    for d in sorted(errors, key=lambda d: (d["error_summary"] or "", d["harness"])):
        kind = ("infrastructure" if d["error_summary"] in infra or (d["observed_log_evidence"] or "").startswith("Package")
                else "time limit" if "timeout" in (d["error_summary"] or "").lower() else "agent / unclear")
        d["_kind"] = kind
    kinds = collections.Counter((d["_kind"], d["harness"]) for d in errors)
    for k in ("infrastructure", "time limit", "agent / unclear"):
        row = {h: kinds[(k, h)] for h in ALL4 if kinds[(k, h)]}
        print(f"  {k:16} {sum(row.values()):2}  {row}")
    setup = [d for d in errors if d["error_summary"] == "Setup exit 100"]
    print("  setup failures (the harness could not install):")
    for d in setup:
        lc = L[(d["task"], d["harness"])]
        luna = f"Luna {lc['crashes']}/{lc['attempts']} attempts crashed" if lc["attempts"] else "Luna has not run it"
        print(f"    {d['task']:28} {d['harness']:15} {luna}")
    conflict = [d for d in errors if (qd.num(d["reward"]) or 0) > 0]
    for d in conflict:
        print(f"  scoring conflict: {d['task']}/{d['harness']} is ERROR but reward={d['reward']}, "
              f"tests {d['tests_passed']}/{d['tests_total']} - counted as no signal in the success table, "
              f"though the work passed")
    out["errors"] = {f"{k}|{h}": v for (k, h), v in kinds.items()}
    out["setup_failures"] = [(d["task"], d["harness"]) for d in setup]
    out["scoring_conflicts"] = [(d["task"], d["harness"]) for d in conflict]

    path = os.path.join(ROOT, "study", "qwen_fails_results.json")
    json.dump(out, open(path, "w"), indent=1, default=float)
    print(f"\nwrote {os.path.relpath(path, ROOT)}")


if __name__ == "__main__":
    main()
